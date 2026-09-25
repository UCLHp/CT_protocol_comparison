import pandas as pd
import xml.etree.ElementTree as ET
import os
import tkinter as tk
from tkinter import filedialog
from openpyxl.styles import PatternFill
from openpyxl import load_workbook
from datetime import datetime


# ========== CONFIG ==========

gray_fill = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")
green_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
orange_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
row_fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")


# ========== STEP 1: Select XML ==========

def _center_window(win, width, height):
    win.update_idletasks()
    sw = win.winfo_screenwidth()
    sh = win.winfo_screenheight()
    x = (sw - width) // 2
    y = (sh - height) // 2
    win.geometry(f"{width}x{height}+{x}+{y}")


def choose_mode_and_files(root):
    choice = {"mode": None}

    dialog = tk.Toplevel(root)
    dialog.title("XML Comparison Tool")
    dialog.resizable(False, False)

    dialog.attributes("-topmost", True)
    _center_window(dialog, 520, 170)

    def set_mode(mode):
        choice["mode"] = mode
        dialog.destroy()

    def cancel():
        choice["mode"] = None
        dialog.destroy()

    dialog.protocol("WM_DELETE_WINDOW", cancel)
    dialog.bind("<Escape>", lambda _e: cancel())

    tk.Label(
        dialog,
        text="Choose mode:"
    ).pack(padx=16, pady=(16, 10))

    frame = tk.Frame(dialog)
    frame.pack(padx=16, pady=(0, 10))

    tk.Button(
        frame,
        text="1 file: Extract",
        width=22,
        command=lambda: set_mode("extract")
    ).grid(row=0, column=0, padx=6, pady=4)

    tk.Button(
        frame,
        text="2 files: Compare",
        width=22,
        command=lambda: set_mode("compare")
    ).grid(row=0, column=1, padx=6, pady=4)

    tk.Button(
        dialog,
        text="Cancel",
        width=12,
        command=cancel
    ).pack(pady=(0, 14))

    dialog.lift()
    dialog.focus_force()

    dialog.after(
        300,
        lambda: dialog.attributes("-topmost", False)
    )

    root.wait_window(dialog)

    mode = choice["mode"]

    if mode is None:
        return []

    if mode == "extract":
        path = filedialog.askopenfilename(
            parent=root,
            title="Select ONE XML file to extract",
            filetypes=[("XML files", "*.xml")]
        )

        return [path] if path else []

    before_path = filedialog.askopenfilename(
        parent=root,
        title="Select BEFORE (older) XML file",
        filetypes=[("XML files", "*.xml")]
    )

    if not before_path:
        return []

    after_path = filedialog.askopenfilename(
        parent=root,
        title="Select AFTER (newer) XML file",
        filetypes=[("XML files", "*.xml")]
    )

    if not after_path:
        return []

    return [before_path, after_path]


def select_output_folder(root):
    return filedialog.askdirectory(
        parent=root,
        title="Select Folder to Save Outputs"
    )


# ========== STEP 2: Parse XML ==========

def parse_privileges(root):
    rows = []
    rows_by_id = {}

    privileges_version = root.findtext("privilegesversion")

    for privilege in root.findall("privilege"):
        row = {}

        row["privilegeid"] = privilege.findtext("privilegeid")
        row["privilegeversion"] = privilege.findtext("privilegeversion")
        row["privilegecategory"] = privilege.findtext("privilegecategory")
        row["privilegegroup"] = privilege.findtext("privilegegroup")
        row["privilegedisplayname"] = privilege.findtext("privilegedisplayname")

        for group in privilege.findall("usergroups/groupcuid"):
            group_name = group.text
            allow_value = group.get("allow")

            row[group_name] = allow_value

        rows.append(row)

        row_id = row["privilegeid"]
        rows_by_id[row_id] = row

    df = pd.DataFrame(rows).fillna("")

    fixed_cols = ["privilegeid", "privilegeversion", "privilegecategory", "privilegegroup", "privilegedisplayname"]

    group_cols = sorted(
        col for col in df.columns
        if col not in fixed_cols
    )

    df = df[fixed_cols + group_cols]
    df = df.sort_values("privilegeid").reset_index(drop=True)

    id_col = "privilegeid"

    return df, rows_by_id, id_col, fixed_cols, privileges_version


def parse_users(root):
    rows = []
    rows_by_id = {}

    for user in root.findall("users/user"):
        row = {}

        row["userid"] = user.findtext("userid")
        row["username"] = user.findtext("username")
        row["isaccountdisabled"] = user.findtext("isaccountdisabled")
        row["groupcuid"] = user.findtext("groupcuid")
        row["provider"] = user.findtext("provider")

        rows.append(row)

        row_id = row["userid"]
        rows_by_id[row_id] = row

    df = pd.DataFrame(rows).fillna("")

    fixed_cols = ["userid", "username", "isaccountdisabled", "groupcuid", "provider"]

    df = df[fixed_cols]
    df = df.sort_values("userid").reset_index(drop=True)

    id_col = "userid"

    return df, rows_by_id, id_col, fixed_cols


def parse_xml_file(xml_path):
    tree = ET.parse(xml_path)
    root = tree.getroot()

    if root.tag == "privileges":
        (
            df,
            rows_by_id,
            id_col,
            fixed_cols,
            file_version
        ) = parse_privileges(root)

        file_type = "privileges"

    elif root.tag == "OSPAccessImport":
        (
            df,
            rows_by_id,
            id_col,
            fixed_cols
        ) = parse_users(root)

        file_type = "users"
        file_version = None

    else:
        raise ValueError(
            f"Unsupported XML type: {root.tag}"
        )

    return df, rows_by_id, file_type, id_col, fixed_cols, file_version


# ========== STEP 3: Compare ==========

def compare_files(
    dict_before,
    dict_after,
    id_col
):
    ids_before = set(dict_before.keys())
    ids_after = set(dict_after.keys())

    removed_ids = ids_before - ids_after
    added_ids = ids_after - ids_before
    common_ids = ids_before & ids_after

    removed_rows = []
    added_rows = []
    changed_rows = []

    for row_id in removed_ids:
        removed_rows.append(
            dict_before[row_id]
        )

    for row_id in added_ids:
        added_rows.append(
            dict_after[row_id]
        )

    all_columns = set()

    for row in dict_before.values():
        all_columns.update(row.keys())

    for row in dict_after.values():
        all_columns.update(row.keys())

    all_columns.remove(id_col)
    all_columns = sorted(all_columns)

    for row_id in common_ids:
        before_row = dict_before[row_id]
        after_row = dict_after[row_id]

        for column in all_columns:
            before_value = before_row.get(column, "")
            after_value = after_row.get(column, "")

            if before_value is None:
                before_value = ""

            if after_value is None:
                after_value = ""

            before_value = str(before_value)
            after_value = str(after_value)

            if before_value != after_value:
                changed_rows.append({
                    id_col: row_id,
                    "Parameter": column,
                    "Before": before_value,
                    "After": after_value
                })

    removed_df = pd.DataFrame(
        removed_rows
    ).fillna("")

    added_df = pd.DataFrame(
        added_rows
    ).fillna("")

    changed_df = pd.DataFrame(
        changed_rows,
        columns=[
            id_col,
            "Parameter",
            "Before",
            "After"
        ]
    )

    return removed_df, added_df, changed_df


def align_columns_for_output(
    df_before,
    df_after,
    fixed_cols
):
    extra_cols = sorted(
        (
            set(df_before.columns)
            | set(df_after.columns)
        )
        - set(fixed_cols)
    )

    all_cols = fixed_cols + extra_cols

    df_before = df_before.reindex(
        columns=all_cols,
        fill_value=""
    )

    df_after = df_after.reindex(
        columns=all_cols,
        fill_value=""
    )

    return df_before, df_after


# ========== STEP 4: Highlight ==========

def highlight_rows(
    filepath,
    highlight_df,
    fill_color,
    id_col
):
    if highlight_df.empty:
        return

    wb = load_workbook(filepath)
    ws = wb.active

    df_excel = pd.read_excel(filepath).fillna("")

    highlight_ids = set(
        highlight_df[id_col]
    )

    for index, row in df_excel.iterrows():
        if row[id_col] in highlight_ids:
            excel_row = index + 2

            for cell in ws[excel_row]:
                cell.fill = fill_color

    wb.save(filepath)


def highlight_changes(
    filepath,
    changed_df,
    id_col
):
    if changed_df.empty:
        return

    wb = load_workbook(filepath)
    ws = wb.active

    df_excel = pd.read_excel(filepath).fillna("")

    changes = {}

    for _, row in changed_df.iterrows():
        row_id = row[id_col]
        parameter = row["Parameter"]

        if row_id == "[FILE]":
            continue

        if row_id not in changes:
            changes[row_id] = []

        changes[row_id].append(parameter)

    for index, row in df_excel.iterrows():
        row_id = row[id_col]

        if row_id in changes:
            excel_row = index + 2

            # Highlight changed row
            for cell in ws[excel_row]:
                cell.fill = row_fill

            # Highlight changed cells
            for parameter in changes[row_id]:
                if parameter in df_excel.columns:
                    column_number = (
                        df_excel.columns.get_loc(parameter)
                        + 1
                    )

                    ws.cell(
                        row=excel_row,
                        column=column_number
                    ).fill = orange_fill

    wb.save(filepath)


def autosize_excel_columns(filepath):
    wb = load_workbook(filepath)
    ws = wb.active

    for column_cells in ws.columns:
        length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        ws.column_dimensions[
            column_cells[0].column_letter
        ].width = length + 2

    wb.save(filepath)


# ========== STEP 5: Main ==========

def main():
    root = tk.Tk()
    root.withdraw()

    xml_files = choose_mode_and_files(root)

    if not xml_files:
        print("Cancelled.")
        raise SystemExit

    save_folder = select_output_folder(root)

    if not save_folder:
        print("Cancelled.")
        raise SystemExit

    # Extract one XML file
    if len(xml_files) == 1:
        (
            df,
            rows_by_id,
            file_type,
            id_col,
            fixed_cols,
            file_version
        ) = parse_xml_file(
            xml_files[0]
        )

        xml_base = os.path.splitext(
            os.path.basename(xml_files[0])
        )[0]

        out_path = os.path.join(
            save_folder,
            f"{xml_base}_summary.xlsx"
        )

        with pd.ExcelWriter(
            out_path,
            engine="openpyxl"
        ) as writer:

            if file_type == "privileges":
                df.to_excel(
                    writer,
                    sheet_name="Privileges",
                    index=False
                )

                file_info = pd.DataFrame({
                    "Setting": ["privilegesversion"],
                    "Value": [file_version]
                })

                file_info.to_excel(
                    writer,
                    sheet_name="File Info",
                    index=False
                )

            else:
                df.to_excel(
                    writer,
                    index=False
                )

        autosize_excel_columns(
            out_path
        )

        print(
            f"Extracted single file to: {out_path}"
        )

        raise SystemExit

    # Compare two XML files
    before_path = xml_files[0]
    after_path = xml_files[1]

    # Create a dated comparison folder
    date_stamp = datetime.now().strftime("%Y%m%d")
    run_number = 1

    while True:
        comparison_folder = os.path.join(
            save_folder,
            f"comparison_{date_stamp}_{run_number}"
        )

        if not os.path.exists(comparison_folder):
            break

        run_number += 1

    os.makedirs(comparison_folder)

    (
        df_before,
        dict_before,
        before_type,
        before_id_col,
        before_fixed_cols,
        before_file_version
    ) = parse_xml_file(
        before_path
    )

    (
        df_after,
        dict_after,
        after_type,
        after_id_col,
        after_fixed_cols,
        after_file_version
    ) = parse_xml_file(
        after_path
    )

    if before_type != after_type:
        raise ValueError(
            "The BEFORE and AFTER XML files are different XML types."
        )

    id_col = before_id_col
    fixed_cols = before_fixed_cols

    (
        removed_df,
        added_df,
        changed_df
    ) = compare_files(
        dict_before,
        dict_after,
        id_col
    )

    # Compare file-level privileges version
    if (
        before_type == "privileges"
        and before_file_version != after_file_version
    ):
        file_change = pd.DataFrame([{
            id_col: "[FILE]",
            "Parameter": "privilegesversion",
            "Before": before_file_version,
            "After": after_file_version
        }])

        changed_df = pd.concat(
            [changed_df, file_change],
            ignore_index=True
        )

    # Create comparison file information
    comparison_info = pd.DataFrame({
        "": [
            "BEFORE",
            "AFTER",
            "Comparison date"
        ],
        "File": [
            os.path.basename(before_path),
            os.path.basename(after_path),
            datetime.now().strftime("%d/%m/%Y")
        ]
    })

    # Save comparison report
    report_path = os.path.join(
        comparison_folder,
        "comparison_report.xlsx"
    )

    with pd.ExcelWriter(
        report_path,
        engine="openpyxl"
    ) as writer:

        comparison_info.to_excel(
            writer,
            sheet_name="File Info",
            index=False
        )

        removed_df.to_excel(
            writer,
            sheet_name="Removed",
            index=False
        )

        added_df.to_excel(
            writer,
            sheet_name="Added",
            index=False
        )

        changed_df.to_excel(
            writer,
            sheet_name="Changed",
            index=False
        )

        for sheet_name in writer.sheets:
            worksheet = writer.sheets[sheet_name]

            for column_cells in worksheet.columns:
                length = max(
                    len(str(cell.value or ""))
                    for cell in column_cells
                )

                worksheet.column_dimensions[
                    column_cells[0].column_letter
                ].width = length + 2

    # Create output filenames
    before_base = os.path.splitext(
        os.path.basename(before_path)
    )[0]

    after_base = os.path.splitext(
        os.path.basename(after_path)
    )[0]

    before_out = os.path.join(
        comparison_folder,
        f"BEFORE_{before_base}_highlighted.xlsx"
    )

    after_out = os.path.join(
        comparison_folder,
        f"AFTER_{after_base}_highlighted.xlsx"
    )

    # Align columns
    df_before, df_after = align_columns_for_output(
        df_before,
        df_after,
        fixed_cols
    )

    # Save BEFORE file
    with pd.ExcelWriter(
        before_out,
        engine="openpyxl"
    ) as writer:

        if before_type == "privileges":
            df_before.to_excel(
                writer,
                sheet_name="Privileges",
                index=False
            )

            file_info = pd.DataFrame({
                "Setting": ["privilegesversion"],
                "Value": [before_file_version]
            })

            file_info.to_excel(
                writer,
                sheet_name="File Info",
                index=False
            )

        else:
            df_before.to_excel(
                writer,
                index=False
            )

    # Save AFTER file
    with pd.ExcelWriter(
        after_out,
        engine="openpyxl"
    ) as writer:

        if after_type == "privileges":
            df_after.to_excel(
                writer,
                sheet_name="Privileges",
                index=False
            )

            file_info = pd.DataFrame({
                "Setting": ["privilegesversion"],
                "Value": [after_file_version]
            })

            file_info.to_excel(
                writer,
                sheet_name="File Info",
                index=False
            )

        else:
            df_after.to_excel(
                writer,
                index=False
            )

    # Highlight removed and added rows
    highlight_rows(
        before_out,
        removed_df,
        gray_fill,
        id_col
    )

    highlight_rows(
        after_out,
        added_df,
        green_fill,
        id_col
    )

    # Highlight changed rows and cells
    highlight_changes(
        before_out,
        changed_df,
        id_col
    )

    highlight_changes(
        after_out,
        changed_df,
        id_col
    )

    autosize_excel_columns(
        before_out
    )

    autosize_excel_columns(
        after_out
    )

    print(
        f"Comparison report saved to: {report_path}"
    )

    print(
        f"Highlighted BEFORE saved to: {before_out}"
    )

    print(
        f"Highlighted AFTER saved to: {after_out}"
    )


if __name__ == "__main__":
    main()